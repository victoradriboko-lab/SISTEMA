#!/usr/bin/env python3
"""Gera manual DOCX editável e PDF a partir do conteúdo compartilhado.
Requer python-docx, reportlab e Pillow. Não consulta dados de produção.
"""
from pathlib import Path
import sys, re
from xml.sax.saxutils import escape
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'docs/manual'))
from conteudo import PAGES
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.colors import HexColor, white
from reportlab.lib.enums import TA_LEFT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

OUT = ROOT / 'docs/manual'
IMAGES = OUT / 'imagens'
NAVY='0A1732'; CYAN='009BB5'; PURPLE='692CDA'; TEXT='263449'; MUTED='58677B'; LIGHT='EFF7FB'
NAME='Manual-do-Usuario-SERVITEC-VSA'


def shade(cell, color):
    el = OxmlElement('w:shd'); el.set(qn('w:fill'), color); cell._tc.get_or_add_tcPr().append(el)

def margins(cell, value=100):
    el = OxmlElement('w:tcMar')
    for name in ('top','left','bottom','right'):
        x=OxmlElement('w:'+name); x.set(qn('w:w'),str(value)); x.set(qn('w:type'),'dxa'); el.append(x)
    cell._tc.get_or_add_tcPr().append(el)

def text_run(p, text, size=None, bold=False, color=TEXT):
    r=p.add_run(text); r.bold=bold; r.font.color.rgb=RGBColor.from_string(color)
    if size: r.font.size=Pt(size)
    return r

def no_split(row):
    x=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(x)

def word():
    d=Document(); s=d.sections[0]
    s.page_width=Cm(21); s.page_height=Cm(29.7)
    s.top_margin=Cm(2); s.bottom_margin=Cm(1.8); s.left_margin=Cm(2.1); s.right_margin=Cm(2.1)
    s.header_distance=Cm(.85); s.footer_distance=Cm(.85); s.different_first_page_header_footer=True
    normal=d.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(10)
    normal.font.color.rgb=RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.08
    for name,size,color in [('Title',36,NAVY),('Heading 1',24,NAVY),('Heading 2',13,CYAN)]:
        st=d.styles[name]; st.font.name='Arial'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
        st.paragraph_format.space_before=Pt(9); st.paragraph_format.space_after=Pt(9)
    for name in ('List Number','List Bullet'):
        d.styles[name].font.name='Arial'; d.styles[name].font.size=Pt(10)
    lang=OxmlElement('w:lang'); lang.set(qn('w:val'),'pt-BR'); normal.element.get_or_add_rPr().append(lang)
    hp=s.header.paragraphs[0]; text_run(hp,'SERVITEC VSA',9,True,NAVY); text_run(hp,'   /   MANUAL DO USUÁRIO',8,color=MUTED)
    border=OxmlElement('w:pBdr'); line=OxmlElement('w:bottom'); line.set(qn('w:val'),'single'); line.set(qn('w:sz'),'10'); line.set(qn('w:color'),CYAN); border.append(line); hp._p.get_or_add_pPr().append(border)
    f=s.footer.paragraphs[0]; text_run(f,'SERVITEC  •  USO E OPERAÇÃO',8,color=MUTED)
    f.add_run('\t'); f.paragraph_format.tab_stops.add_tab_stop(Cm(14.7))
    text_run(f,'PÁGINA ',8,color=MUTED)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
    # Capa com um bloco de marca, título editável e hierarquia clara.
    d.add_paragraph('MANUAL OFICIAL DE UTILIZAÇÃO',style='Subtitle')
    d.add_paragraph('')
    t=d.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    cell=t.cell(0,0); shade(cell,NAVY); margins(cell,380)
    p=cell.paragraphs[0]; text_run(p,'SISTEMA INFORMÁTICA',12,True,'62DCEB')
    p=cell.add_paragraph(); text_run(p,'SERVITEC',34,True,'FFFFFF')
    p=cell.add_paragraph(); text_run(p,'VSA',50,True,'62DCEB')
    p=cell.add_paragraph(); text_run(p,'TECNOLOGIA  /  ATENDIMENTO  /  GESTÃO',9,False,'DCE8F6')
    d.add_paragraph('')
    d.add_paragraph('Manual do usuário',style='Title')
    p=d.add_paragraph(); text_run(p,'Da rotina de atendimento à visão do cliente.',16,color=MUTED)
    d.add_paragraph('Guia prático para administradores, operadores e clientes.\nChamados • Inventário • Monitoramento • Relatórios • Backups')
    d.add_paragraph('')
    p=d.add_paragraph(); text_run(p,'EDIÇÃO 1.2',10,True,CYAN); text_run(p,'   |   OUTUBRO DE 2026',10,color=MUTED)
    d.add_paragraph('Documento editável em Word.\nTelas ilustradas com dados de demonstração.')
    for title,kicker,blocks in PAGES:
        d.add_page_break()
        p=d.add_paragraph(); text_run(p,kicker,9,True,CYAN)
        d.add_paragraph(title,style='Heading 1')
        for item in blocks:
            typ=item[0]
            if typ=='p': d.add_paragraph(item[1])
            elif typ=='h': d.add_paragraph(item[1],style='Heading 2')
            elif typ=='steps':
                for i,txt in enumerate(item[1],1):
                    p=d.add_paragraph(); p.paragraph_format.left_indent=Cm(.1)
                    text_run(p,f'{i:02d}  ',10,True,CYAN); text_run(p,txt)
            elif typ=='code':
                t=d.add_table(rows=1,cols=1); c=t.cell(0,0); shade(c,'F1F4F8'); margins(c,130)
                r=c.paragraphs[0].add_run(item[1]); r.font.name='Consolas'; r.font.size=Pt(8.5)
                c.paragraphs[0].paragraph_format.space_after=Pt(0)
                no_split(t.rows[0]); d.add_paragraph().paragraph_format.space_after=Pt(0)
            elif typ=='note':
                t=d.add_table(rows=1,cols=1); c=t.cell(0,0); shade(c,LIGHT); margins(c,135)
                p=c.paragraphs[0]; text_run(p,item[1].upper(),9,True,CYAN)
                p=c.add_paragraph(); p.paragraph_format.space_after=Pt(0); text_run(p,item[2],9)
                no_split(t.rows[0])
                d.add_paragraph().paragraph_format.space_after=Pt(0)
            elif typ=='table':
                t=d.add_table(rows=0,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
                t.columns[0].width=Cm(5.1); t.columns[1].width=Cm(11.7)
                for idx,row in enumerate(item[1]):
                    cells=t.add_row().cells
                    for j,txt in enumerate(row):
                        c=cells[j]; margins(c,100); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        shade(c,NAVY if idx==0 else ('F1F5FA' if idx%2 else 'FFFFFF'))
                        p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.05
                        text_run(p,txt,9.2,idx==0,'FFFFFF' if idx==0 else TEXT)
                    no_split(t.rows[-1])
                    if idx==0:
                        x=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(x)
                d.add_paragraph().paragraph_format.space_after=Pt(0)
            elif typ=='image':
                p=d.add_paragraph(); p.paragraph_format.space_after=Pt(3)
                p.add_run().add_picture(str(IMAGES/item[1]),width=Cm(16.5))
                p=d.add_paragraph(); p.paragraph_format.space_after=Pt(9); text_run(p,item[2],8,color=MUTED)
    d.core_properties.title='Manual do usuário — Sistema Informática SERVITEC VSA'
    d.core_properties.subject='Guia de utilização da plataforma'
    d.core_properties.author='SERVITEC VSA'; d.core_properties.keywords='manual, chamados, inventário, portal, backups'
    d.core_properties.comments='Imagens e registros demonstrativos. Edição 1.2; versão 0.4.0 com contratos, SLA e segurança.'
    d.save(OUT/(NAME+'.docx'))


def pdf():
    fontdir=Path('/usr/share/fonts/truetype/liberation')
    for label,file in [('Arial','LiberationSans-Regular.ttf'),('Arial-Bold','LiberationSans-Bold.ttf'),('Mono','LiberationMono-Regular.ttf')]:
        pdfmetrics.registerFont(TTFont(label,str(fontdir/file)))
    pdfmetrics.registerFontFamily('Arial',normal='Arial',bold='Arial-Bold',italic='Arial',boldItalic='Arial-Bold')
    styles={
        'p':ParagraphStyle('Body',fontName='Arial',fontSize=10,leading=12.8,textColor=HexColor('#'+TEXT),spaceAfter=7),
        'small':ParagraphStyle('Small',fontName='Arial',fontSize=8,leading=10,textColor=HexColor('#'+MUTED),spaceAfter=9),
        'h':ParagraphStyle('Sub',fontName='Arial-Bold',fontSize=13,leading=16,textColor=HexColor('#'+CYAN),spaceBefore=10,spaceAfter=7),
        'title':ParagraphStyle('Title',fontName='Arial-Bold',fontSize=24,leading=28,textColor=HexColor('#'+NAVY),spaceAfter=14),
        'kicker':ParagraphStyle('Kicker',fontName='Arial-Bold',fontSize=9,leading=12,textColor=HexColor('#'+CYAN),spaceAfter=8),
        'cell':ParagraphStyle('Cell',fontName='Arial',fontSize=9.2,leading=11.6,textColor=HexColor('#'+TEXT)),
        'head':ParagraphStyle('Head',fontName='Arial-Bold',fontSize=9.2,leading=11.6,textColor=white),
        'code':ParagraphStyle('Code',fontName='Mono',fontSize=8.5,leading=12,textColor=HexColor('#'+TEXT),backColor=HexColor('#F1F4F8'),borderPadding=10,spaceBefore=5,spaceAfter=13),
    }
    para=lambda txt,style='p':Paragraph(escape(txt).replace('\n','<br/>'),styles[style])
    width=476.2
    def onpage(canvas,doc):
        canvas.saveState()
        if doc.page>1:
            canvas.setFont('Arial-Bold',9);canvas.setFillColor(HexColor('#'+NAVY));canvas.drawString(59.5,811,'SERVITEC VSA')
            canvas.setFont('Arial',8);canvas.setFillColor(HexColor('#'+MUTED));canvas.drawString(145,811,'/  MANUAL DO USUÁRIO')
            canvas.setStrokeColor(HexColor('#'+CYAN));canvas.setLineWidth(1);canvas.line(59.5,803,535.7,803)
            canvas.setFont('Arial',8);canvas.drawString(59.5,27,'SERVITEC  •  USO E OPERAÇÃO');canvas.drawRightString(535.7,27,f'PÁGINA {doc.page:02d}')
        canvas.restoreState()
    story=[Spacer(1,12),para('MANUAL OFICIAL DE UTILIZAÇÃO','kicker'),Spacer(1,32)]
    coverstyles=[ParagraphStyle('BrandSmall',fontName='Arial-Bold',fontSize=12,leading=20,textColor=HexColor('#62DCEB')),ParagraphStyle('Brand',fontName='Arial-Bold',fontSize=40,leading=52,textColor=white),ParagraphStyle('VSA',fontName='Arial-Bold',fontSize=58,leading=68,textColor=HexColor('#62DCEB'))]
    cover=Table([[Paragraph('SISTEMA INFORMÁTICA',coverstyles[0])],[Paragraph('SERVITEC',coverstyles[1])],[Paragraph('VSA',coverstyles[2])],[Paragraph('TECNOLOGIA  /  ATENDIMENTO  /  GESTÃO',styles['head'])]],colWidths=[width])
    cover.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),HexColor('#'+NAVY)),('LEFTPADDING',(0,0),(-1,-1),23),('TOPPADDING',(0,0),(-1,0),25),('BOTTOMPADDING',(0,-1),(-1,-1),25)]))
    story += [cover,Spacer(1,35),Paragraph('Manual do usuário',ParagraphStyle('CoverTitle',parent=styles['title'],fontSize=34,leading=40)),Spacer(1,12),para('Da rotina de atendimento à visão do cliente.','h'),Spacer(1,13),para('Guia prático para administradores, operadores e clientes.\nChamados • Inventário • Monitoramento • Relatórios • Backups'),Spacer(1,24),para('EDIÇÃO 1.2   |   OUTUBRO DE 2026','kicker'),para('Telas ilustradas com dados de demonstração.','small')]
    for title,kicker,blocks in PAGES:
        story += [PageBreak(),para(kicker,'kicker'),para(title,'title')]
        for item in blocks:
            typ=item[0]
            if typ in ('p','h','code'): story.append(para(item[1],typ))
            elif typ=='steps':
                for i,txt in enumerate(item[1],1):
                    story.append(Paragraph(f'<font color="#{CYAN}"><b>{i:02d}</b></font>  '+escape(txt),styles['p']))
            elif typ=='image':
                from PIL import Image as PILImage
                with PILImage.open(IMAGES/item[1]) as im: iw,ih=im.size
                story += [Image(str(IMAGES/item[1]),width=467.7,height=467.7*ih/iw),Spacer(1,5),para(item[2],'small')]
            elif typ=='note':
                c=[para(item[1].upper(),'kicker'),para(item[2],'p')]
                t=Table([[c]],colWidths=[width]);t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),HexColor('#'+LIGHT)),('BOX',(0,0),(-1,-1),.3,HexColor('#D1E8F2')),('LEFTPADDING',(0,0),(-1,-1),11),('RIGHTPADDING',(0,0),(-1,-1),11),('TOPPADDING',(0,0),(-1,-1),10),('BOTTOMPADDING',(0,0),(-1,-1),5)]));story.extend([t,Spacer(1,7)])
            elif typ=='table':
                data=[[para(txt,'head' if idx==0 else 'cell') for txt in row] for idx,row in enumerate(item[1])]
                t=Table(data,colWidths=[144,332.2],repeatRows=1,hAlign='LEFT')
                ts=[('BACKGROUND',(0,0),(-1,0),HexColor('#'+NAVY)),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]
                for idx in range(1,len(data)):
                    if idx%2: ts.append(('BACKGROUND',(0,idx),(-1,idx),HexColor('#F1F5FA')))
                t.setStyle(TableStyle(ts));story.extend([t,Spacer(1,9)])
    doc=SimpleDocTemplate(str(OUT/(NAME+'.pdf')),pagesize=(595.2756,841.8898),leftMargin=59.5,rightMargin=59.5,topMargin=56.7,bottomMargin=51,title='Manual do usuário — SERVITEC VSA',author='SERVITEC VSA')
    doc.build(story,onFirstPage=onpage,onLaterPages=onpage)

if __name__=='__main__':
    word();pdf();print(f'Manual Word e PDF gerados em {OUT}')
